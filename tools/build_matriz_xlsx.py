#!/usr/bin/env python3
"""Genera matrices/matriz_riesgos.xlsx (BORRADOR inicial para que el equipo lo refine).

Hojas: «Inicial», «Residual» y «Escala». Misma escala y metodología en ambas matrices
(probabilidad × impacto, 1–5; umbrales 1–4 bajo, 5–9 moderado, 10–16 alto, 17–25 crítico).

Los valores de P e I son estimaciones iniciales: deben justificarse y ajustarse con la
evidencia del laboratorio. Una vez editado el Excel a mano, NO vuelva a ejecutar este
script (sobrescribiría los cambios): sirve solo para regenerar la plantilla.

Uso:  python tools/build_matriz_xlsx.py   (requiere openpyxl; luego recalcule con LibreOffice si desea valores en caché)
"""
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

OUT = Path(__file__).resolve().parents[1] / "matrices" / "matriz_riesgos.xlsx"
F = "Arial"
thin = Side(style="thin", color="7F7F7F")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
LEVEL_COLORS = {"Crítico": ("C00000", "FFFFFF"), "Alto": ("F4B183", "000000"), "Moderado": ("FFE699", "000000"), "Bajo": ("C6E0B4", "000000")}

# id, activo, amenaza, vulnerabilidad, escenario, impacto, probabilidad, controles_existentes, justificación
INICIAL = [
 ("R-01", "Cuenta de usuario ficticia «usuario.demo» en la aplicación de autenticación (formulario /login)",
  "Atacante que prueba contraseñas de forma repetida (fuerza bruta / adivinación).",
  "La aplicación acepta intentos ilimitados de inicio de sesión y no distingue intentos consecutivos fallidos.",
  "Un atacante automatiza intentos contra /login hasta adivinar la contraseña y accede a la cuenta; consecuencia: acceso no autorizado a los servicios del portal.",
  3, 4, "Contraseña almacenada con hash (PBKDF2) y mensaje de error genérico. Sin bloqueo ni registro de auditoría.",
  "P=4: no hay límite de intentos y la herramienta de ataque es trivial. I=3: la cuenta es ficticia pero representa una cuenta de funcionario con acceso a servicios."),
 ("R-02", "Cuenta de usuario ficticia (credenciales del funcionario) expuesta a un mensaje con QR / enlace corto",
  "Phishing: mensaje simulado con QR o enlace acortado que lleva a un formulario falso.",
  "El usuario no puede ver el destino real de un QR o enlace corto; no hay segunda verificación (MFA); poca concienciación.",
  "El usuario escanea el QR, ingresa sus credenciales en el formulario y el atacante las reutiliza; consecuencia: suplantación de identidad y acceso a datos.",
  4, 4, "Ninguno (sin MFA ni filtrado de enlaces/QR en el laboratorio).",
  "P=4: el phishing por QR/enlace es frecuente y efectivo (ver casos). I=4: credenciales válidas permiten suplantar al usuario."),
 ("R-03", "Registros de acceso: auth.log de la aplicación y access.log de Apache/NGINX",
  "Incidente sin detección ni reconstrucción posible (falta de trazabilidad).",
  "No existe registro de eventos de autenticación en la aplicación; los logs del servidor no muestran usuario ni estado de bloqueo.",
  "Se producen intentos fallidos repetidos y nadie puede reconstruir la secuencia ni saber cuándo ni desde dónde; consecuencia: respuesta tardía y falta de evidencia.",
  3, 4, "access.log del servidor web (IP, hora, método, código HTTP). Sin registro a nivel de aplicación.",
  "P=4: sin auditoría el evento ocurre sin ser detectado. I=3: afecta la capacidad de respuesta y de rendir cuentas."),
 ("R-04", "Servidor web Apache HTTP Server (proxy inverso de la aplicación)",
  "Reconocimiento y explotación de fallas conocidas del software expuesto.",
  "Encabezados y mensajes revelan la tecnología interna (p. ej. «Server: gunicorn»); configuración por defecto.",
  "Un atacante identifica versión/tecnología y prueba exploits conocidos; consecuencia: compromiso del servidor o indisponibilidad del portal.",
  3, 3, "Ninguno (configuración por defecto).",
  "P=3: la exposición es posible pero requiere una falla explotable. I=3: afecta la disponibilidad del portal de pruebas."),
 ("R-05", "Servidor web NGINX (proxy inverso de la aplicación)",
  "Reconocimiento y explotación de fallas conocidas del software expuesto.",
  "Encabezados y mensajes revelan la tecnología interna; configuración por defecto.",
  "Igual que R-04 sobre NGINX; se mantiene como riesgo separado para comparar ambos servidores.",
  3, 3, "Ninguno (configuración por defecto).",
  "Misma justificación que R-04, aplicada a NGINX."),
 ("R-06", "Tráfico HTTP de autenticación (credenciales en tránsito entre cliente y servidor)",
  "Interceptación del tráfico en la red (sniffing / intermediario).",
  "El laboratorio usa HTTP sin cifrado (TLS); las credenciales viajan en claro dentro de la red del laboratorio.",
  "Un equipo en la misma red captura el POST /login y obtiene usuario y contraseña; consecuencia: robo de credenciales.",
  4, 3, "Red aislada del laboratorio (control de entorno, no de la aplicación).",
  "P=3: requiere estar en la misma red. I=4: expone directamente las credenciales."),
]

# id, control, iso, nist, evidencia, P, I, permanece
RESIDUAL = [
 ("R-01", "Bloqueo temporal de la cuenta tras 3 intentos fallidos consecutivos (120 s) + mensaje informativo + registro de activación y fin del bloqueo.",
  "Tratamiento del riesgo (modificar): selección de control que reduce la probabilidad; aceptación del riesgo residual por el responsable.",
  "Proteger (PR): PR.AA Gestión de identidades, autenticación y control de acceso. Detectar (DE): DE.CM Monitoreo continuo.",
  "PENDIENTE: enlazar resultados de scripts/run_scenario.py (S1–S3) y extractos de auth.log (LOCKOUT_START/LOCKOUT_END).",
  2, 3, "Un ataque lento o distribuido (varias IP) puede quedar bajo el umbral; el bloqueo puede usarse para negar el servicio a la cuenta legítima. Acciones: límite por IP, MFA, alertas."),
 ("R-02", "Control parcial en el laboratorio: registro del evento LINK_VISIT (campaña) y bloqueo de la cuenta; recomendado: MFA y capacitación con simulacros.",
  "Tratamiento del riesgo (modificar / compartir con el usuario): controles técnicos y de concienciación.",
  "Proteger (PR): PR.AT Concienciación y capacitación; PR.AA Autenticación. Detectar (DE): DE.AE Análisis de eventos adversos.",
  "PENDIENTE: capturas del mensaje simulado, del enlace /r/qr1 y del evento LINK_VISIT en auth.log.",
  3, 4, "Las credenciales capturadas siguen siendo válidas hasta el cambio; el bloqueo no impide el uso inmediato de una contraseña correcta. Acciones: MFA, capacitación periódica, filtrado de QR/enlaces."),
 ("R-03", "Registro de auditoría de autenticación (JSON Lines) con fecha/hora, usuario ficticio, resultado, IP y estado de bloqueo, sin contraseñas ni tokens.",
  "Monitoreo y revisión del riesgo; tratamiento (modificar): control de detección y trazabilidad.",
  "Detectar (DE): DE.CM Monitoreo continuo; DE.AE Análisis de eventos adversos.",
  "PENDIENTE: extracto de auth.log y salida de scripts/analyze_logs.py que reconstruya la secuencia.",
  2, 3, "Los registros están en el mismo host y sin protección de integridad. Acciones: envío a un servidor central/SIEM, retención y control de acceso a los logs."),
 ("R-04", "Endurecimiento de la configuración: ocultar la tecnología interna (encabezado Server), ServerTokens Prod, nosniff, TRACE deshabilitado.",
  "Tratamiento del riesgo (modificar): reducción de la superficie de exposición.",
  "Proteger (PR): PR.PS Seguridad de la plataforma.",
  "PENDIENTE: captura de los encabezados de respuesta antes/después (curl -I o ZAP).",
  2, 3, "Siguen existiendo fallas desconocidas o de versión. Acciones: parches periódicos, TLS, escaneo de vulnerabilidades."),
 ("R-05", "Endurecimiento equivalente en NGINX: server_tokens off, proxy_hide_header Server, nosniff.",
  "Tratamiento del riesgo (modificar): reducción de la superficie de exposición.",
  "Proteger (PR): PR.PS Seguridad de la plataforma.",
  "PENDIENTE: captura de los encabezados de respuesta antes/después.",
  2, 3, "Igual que R-04. Acciones: parches periódicos, TLS, escaneo de vulnerabilidades."),
 ("R-06", "Ninguno implementado en el laboratorio (limitación). Recomendado: HTTPS con TLS y HSTS.",
  "Tratamiento del riesgo: se documenta como limitación; el riesgo se acepta temporalmente solo por tratarse de un laboratorio aislado.",
  "Proteger (PR): PR.DS Seguridad de los datos (en tránsito).",
  "PENDIENTE: captura de ZAP mostrando el POST /login en claro (campos enmascarados en el repositorio).",
  3, 4, "El riesgo permanece igual: un control no implementado no reduce el riesgo. Acción: habilitar TLS en Apache y NGINX."),
]

SC = "'Escala'!$A$12:$A$15"
SN = "'Escala'!$B$12:$B$15"


def hdr(ws, row, labels, widths):
    for c, (h, w) in enumerate(zip(labels, widths), 1):
        cell = ws.cell(row, c, h)
        cell.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[row].height = 36


def body(ws, row, vals, widths, center=()):
    lines = 1
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row, c, v)
        cell.font = Font(name=F, size=9, bold=(c in center and c == 1))
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if c in center else "left")
        cell.border = BORDER
        if isinstance(v, str) and not v.startswith("="):
            n = sum(max(1, math.ceil(len(p) / (widths[c - 1] * 1.25))) for p in v.split("\n"))
            lines = max(lines, n)
    ws.row_dimensions[row].height = min(300, lines * 12 + 10)


def level_cf(ws, col, first, last):
    for name, (bg, fg) in LEVEL_COLORS.items():
        ws.conditional_formatting.add(
            f"{col}{first}:{col}{last}",
            FormulaRule(formula=[f'{col}{first}="{name}"'], fill=PatternFill("solid", bgColor=bg, fgColor=bg),
                        font=Font(name=F, bold=True, color=fg), stopIfTrue=True))


wb = Workbook()
ini = wb.active
ini.title = "Inicial"
res = wb.create_sheet("Residual")
esc = wb.create_sheet("Escala")

# ---------------------------------------------------------------- Inicial
ini["A1"] = "Matriz de riesgos INICIAL — Laboratorio de phishing y monitoreo (BORRADOR)"
ini["A1"].font = Font(name=F, size=14, bold=True, color="1F3864")
ini["A2"] = ("Valores de P e I = estimación inicial del equipo; justificarlos y ajustarlos con la evidencia. "
             "Nivel = Probabilidad × Impacto (escala 1–5). Umbrales en la hoja «Escala».")
ini["A2"].font = Font(name=F, size=9, italic=True, color="404040")
H1 = ["ID", "Activo afectado (concreto)", "Amenaza", "Vulnerabilidad", "Escenario de riesgo", "Impacto (1–5)", "Probabilidad (1–5)",
      "Nivel de riesgo inherente (P×I)", "Clasificación", "Controles existentes", "Justificación de P e I"]
W1 = [7, 34, 30, 34, 42, 10, 12, 14, 13, 34, 44]
hdr(ini, 4, H1, W1)
for k, r in enumerate(INICIAL):
    row = 5 + k
    body(ini, row, [r[0], r[1], r[2], r[3], r[4], r[5], r[6], f"=F{row}*G{row}", f"=LOOKUP(H{row},{SC},{SN})", r[7], r[8]], W1, center=(1, 6, 7, 8, 9))
last = 4 + len(INICIAL)
level_cf(ini, "I", 5, last)
ini.freeze_panes = "C5"

# ---------------------------------------------------------------- Residual
res["A1"] = "Matriz de riesgos RESIDUAL — Laboratorio de phishing y monitoreo (BORRADOR)"
res["A1"].font = Font(name=F, size=14, bold=True, color="1F3864")
res["A2"] = ("Misma escala y metodología que la matriz inicial. P e I se vuelven a estimar considerando la eficacia esperada de los controles; "
             "reemplazar «PENDIENTE» con la evidencia real del laboratorio. Un control no hace desaparecer el riesgo automáticamente.")
res["A2"].font = Font(name=F, size=9, italic=True, color="404040")
H2 = ["ID", "Activo afectado", "Control implementado", "Relación con ISO/IEC 27005", "Función / categoría NIST CSF 2.0",
      "Evidencia de funcionamiento", "Probabilidad residual (1–5)", "Impacto residual (1–5)", "Nivel residual (P×I)", "Clasificación residual",
      "Nivel inicial", "Variación (residual − inicial)", "Riesgo que permanece y acciones adicionales"]
W2 = [7, 28, 40, 36, 34, 36, 12, 12, 12, 13, 10, 13, 44]
hdr(res, 4, H2, W2)
for k, r in enumerate(RESIDUAL):
    row = 5 + k
    body(res, row, [r[0], f"=Inicial!B{row}", r[1], r[2], r[3], r[4], r[5], r[6], f"=G{row}*H{row}", f"=LOOKUP(I{row},{SC},{SN})",
                    f"=Inicial!H{row}", f"=I{row}-K{row}", r[7]], W2, center=(1, 7, 8, 9, 10, 11, 12))
level_cf(res, "J", 5, last)
res.freeze_panes = "C5"

# ---------------------------------------------------------------- Escala
esc.column_dimensions["A"].width = 16
esc.column_dimensions["B"].width = 18
esc.column_dimensions["C"].width = 18
esc["A1"] = "Escala de valoración (definida en la práctica)"
esc["A1"].font = Font(name=F, size=14, bold=True, color="1F3864")


def eh(row, labels):
    for c, h in enumerate(labels, 1):
        cell = esc.cell(row, c, h)
        cell.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER


def er(row, vals):
    for c, v in enumerate(vals, 1):
        cell = esc.cell(row, c, v)
        cell.font = Font(name=F, size=10)
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER


eh(4, ["Valor", "Probabilidad", "Impacto"])
for i, (p, im) in enumerate([("Rara", "Insignificante"), ("Improbable", "Menor"), ("Posible", "Moderado"), ("Probable", "Mayor"), ("Casi segura", "Crítico")], 1):
    er(4 + i, [i, p, im])
esc["A10"] = "Nivel de riesgo R = P × I"
esc["A10"].font = Font(name=F, size=10, bold=True)
eh(11, ["Desde (R ≥)", "Nivel", "Rango"])
for i, (lo, n, rg) in enumerate([(1, "Bajo", "1–4"), (5, "Moderado", "5–9"), (10, "Alto", "10–16"), (17, "Crítico", "17–25")]):
    er(12 + i, [lo, n, rg])
    bg, fg = LEVEL_COLORS[n]
    esc.cell(12 + i, 2).fill = PatternFill("solid", fgColor=bg)
    esc.cell(12 + i, 2).font = Font(name=F, size=10, bold=True, color=fg)
esc["A18"] = "Estos umbrales son una convención definida para la práctica."
esc["A18"].font = Font(name=F, size=9, italic=True)

for ws in (ini, res):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 8
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "4:4"

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.properties.creator = "Marcelo Acuña"
wb.save(OUT)
print("Generado", OUT)
