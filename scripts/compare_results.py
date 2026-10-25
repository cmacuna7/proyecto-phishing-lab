#!/usr/bin/env python3
"""Genera la tabla comparativa antes/después y Apache/NGINX (informe 3.5.4).

Lee los JSON que produce run_scenario.py y, si existen, cuenta los eventos de
auditoría de cada corrida (por run_id) para verificar que se registraron bien.
"""
from __future__ import annotations

import argparse
import glob
import json
import sys
from pathlib import Path


def audit_count(audit_dir: Path, target: str, mode: str, run_id: str) -> str:
    # Una corrida guarda logs/<modo>/app-<servidor>/; el laboratorio en vivo, logs/app-<servidor>/.
    p = audit_dir / mode / f"app-{target}" / "auth.log"
    if not p.exists():
        p = audit_dir / f"app-{target}" / "auth.log"
    if not p.exists():
        return "0 (sin auditoría)" if mode == "baseline" else "n/d"
    n = 0
    for line in p.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        e = json.loads(line)
        if e.get("run_id") == run_id and e["event"] in ("LOGIN_SUCCESS", "LOGIN_FAILED", "LOGIN_BLOCKED"):
            n += 1
    return str(n)


def risk_level(run):
    """Nivel R-01 de matrices/matriz_riesgos.xlsx: inicial (baseline) o residual (hardened). Necesita openpyxl."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(Path(__file__).resolve().parents[1] / "matrices" / "matriz_riesgos.xlsx", data_only=True)
        if run["meta"]["mode"] == "baseline":
            ws = wb["Inicial"]; v, c = ws["H5"].value, ws["I5"].value
        else:
            ws = wb["Residual"]; v, c = ws["I5"].value, ws["J5"].value
        return f"{v} ({c})" if v is not None else "(abrir la matriz y recalcular)"
    except Exception:
        return "(ver matriz)"


def main():
    sys.stdout.reconfigure(encoding="utf-8")  # la consola de Windows (cp1252) no imprime → ni ✔
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--results", default="evidence/corridas/*/resultados/*.json")
    ap.add_argument("--logs-dir", default="", help="Por defecto, la carpeta logs/ de cada corrida")
    ap.add_argument("--out", default="evidence/corridas/comparacion.md")
    a = ap.parse_args()

    paths = sorted(glob.glob(a.results))
    runs = [json.loads(Path(p).read_text(encoding="utf-8"))
            | {"corrida": Path(p).parent.parent.name, "logs": Path(a.logs_dir or Path(p).parent.parent / "logs")}
            for p in paths]
    if not runs:
        raise SystemExit("No hay resultados. Ejecute primero scripts/run_scenario.py")
    runs.sort(key=lambda r: (r["corrida"], r["meta"]["mode"] != "baseline", r["meta"]["target"], r["meta"]["started"]))

    def get(r, step_prefix):
        return next((x for x in r["attempts"] if x["step"].startswith(step_prefix)), None)

    # Si se comparan varias corridas (personas), la columna dice de cuál es cada una.
    several = len({r["corrida"] for r in runs}) > 1
    cols = [f"{r['meta']['target']} · {r['meta']['mode']}" + (f"<br>{r['corrida']}" if several else "") for r in runs]
    rows = [
        ("Intentos fallidos ejecutados", [str(sum(x["password_kind"] == "wrong" for x in r["attempts"])) for r in runs]),
        ("Intentos registrados por la app (auditoría)",
         [audit_count(r["logs"], r["meta"]["target"], r["meta"]["mode"], r["meta"]["run_id"]) for r in runs]),
        ("Bloqueo activado al 3.er fallo consecutivo",
         ["Sí" if get(r, "S1")["lock_message"] is not None and any(x["lock_message"] and x["attempt_no"] == 3 for x in r["attempts"]) else "No" for r in runs]),
        ("4.º intento fallido (HTTP)", [str(get(r, "S1b")["http_status"]) for r in runs]),
        ("Credenciales correctas durante el bloqueo (HTTP)", [str(get(r, "S2")["http_status"]) for r in runs]),
        ("Credenciales correctas tras finalizar el bloqueo (HTTP)", [str(get(r, "S3")["http_status"]) for r in runs]),
        ("Todos los pasos coinciden con lo esperado", ["Sí" if all(x["passed"] for x in r["attempts"]) else "No" for r in runs]),
        ("Nivel de riesgo estimado de fuerza bruta (R-01)", [risk_level(r) for r in runs]),
    ]
    out = ["# Comparación antes y después / Apache vs NGINX", "",
           "| Criterio | " + " | ".join(cols) + " |", "|---|" + "---|" * len(cols)]
    out += [f"| {name} | " + " | ".join(vals) + " |" for name, vals in rows]
    out += ["", "Notas: 'baseline' = sin controles; 'hardened' = con bloqueo y registro de auditoría. "
            "Las diferencias entre Apache y NGINX deben explicarse en el informe a partir de los registros."]
    text = "\n".join(out) + "\n"
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(text, encoding="utf-8")
    print(text)


if __name__ == "__main__":
    main()
